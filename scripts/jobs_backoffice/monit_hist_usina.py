import pandas as pd
from datetime import datetime
import datetime as dt
import sys
import os
import yaml
import re
from pathlib import Path
from colorama import Fore, Style
import warnings
from openpyxl import load_workbook

warnings.filterwarnings("ignore")

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from utils.src_ingestion import initialize_connections
from jobs_backoffice.rateio import LEITORES, validar_soma_rateio

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

print(f"{Fore.MAGENTA}\n=== ETAPA 1: Inicialização ==={Style.RESET_ALL}")


# ============================================================
# CONFIGURAÇÕES
# ============================================================

distribuidora_escolhida = "EQUATORIAL GO"
# opções:
# "CPFL PAULISTA"
# "EQUATORIAL GO"   

data_hoje = pd.Timestamp(datetime.today().date())

mes_pasta_rateio = "07.2026"
allocation_month = "2026-07"

# ------------------------------------------------------------
# Raiz da pasta "Acompanhamento - Operação GD" no OneDrive.
# É a ÚNICA parte que muda entre máquinas — ajuste só aqui ao rodar em outra
# máquina; folder_hist e a pasta de rateio são derivados dela.
# ------------------------------------------------------------
# base_gd = r"C:\Users\livia.medeiros\Serena\Comercializadora - Comercializadora\OMG\03 - BackOffice\13 - Geração Distribuída\Acompanhamento - Operação GD"
base_gd = r"C:\Users\vinicius.azevedo\Serena\Comercializadora - OMG\03 - BackOffice\13 - Geração Distribuída\Acompanhamento - Operação GD"

folder_hist = os.path.join(base_gd, "inputs_manuais", "hist_rateio", "_processado")
base_rateio = os.path.join(base_gd, "7. Rateio")

# A pasta de cada distribuidora dentro de "7. Rateio" é, por padrão, o próprio
# nome dela (distribuidora_escolhida, linha 28). Só registre aqui se a pasta
# tiver nome diferente do valor de distribuidora_escolhida.
pastas_especiais = {
    # "EXEMPLO DIST": "Nome Diferente Da Pasta",
}


# ============================================================
# FUNÇÕES GERAIS
# ============================================================

def converter_mes_envio(valor):
    if pd.isna(valor):
        return pd.NaT

    if isinstance(valor, (pd.Timestamp, datetime, dt.date)):
        return pd.to_datetime(valor, errors="coerce").to_period("M").to_timestamp()

    valor_str = str(valor).strip()

    data = pd.to_datetime("01/" + valor_str, format="%d/%m/%Y", errors="coerce")

    if pd.isna(data):
        data = pd.to_datetime(valor_str, errors="coerce", dayfirst=True)

    if pd.isna(data):
        return pd.NaT

    return data.to_period("M").to_timestamp()


mes_envio = converter_mes_envio("04/2026")


def get_latest_excel_file(folder_path: str, allocation_month: str):
    folder = Path(folder_path)
    month_str = dt.datetime.strptime(allocation_month, "%Y-%m").strftime("%Y_%m_%d")

    if not folder.exists():
        raise FileNotFoundError(f"Pasta não encontrada: {folder_path}")

    files = {
        f.name: f
        for f in folder.iterdir()
        if f.is_file()
        and f.suffix.lower() in [".xlsx", ".xlsm"]
        and not f.name.startswith("~$")
        and re.search(month_str, f.name)
    }

    if not files:
        raise FileNotFoundError(
            f"Nenhum arquivo Excel encontrado para {month_str} em: {folder_path}"
        )

    latest_filename = sorted(files.keys())[-1]
    latest_file = files.get(latest_filename)

    print(f"Arquivo Excel encontrado: {latest_file}")
    return latest_file


def padronizar_chaves_para_merge(df):
    if "Distribuidora" in df.columns:
        df["Distribuidora"] = (
            df["Distribuidora"]
            .astype(str)
            .str.strip()
            .str.upper()
        )

    if "DISTRIBUIDORA" in df.columns:
        df["DISTRIBUIDORA"] = (
            df["DISTRIBUIDORA"]
            .astype(str)
            .str.strip()
            .str.upper()
        )

    for col in [
        "UG", "UC", "CC", "Usina", "CC/UG", "Acrônimo",
        "USINA", "ACRONIMO", "DISTRIBUIDORA"
    ]:
        if col in df.columns:
            df[col] = (
                df[col]
                .astype(str)
                .str.replace(".0", "", regex=False)
                .str.strip()
            )

            df[col] = df[col].replace({
                "nan": pd.NA,
                "None": pd.NA,
                "<NA>": pd.NA,
                "": pd.NA
            })

    return df


def padronizar_datas_rateio(df):
    for col in ["Data de Envio", "DATA_DE_ENVIO"]:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce", dayfirst=True)

    for col in ["Mês de Envio", "MES_DE_ENVIO"]:
        if col in df.columns:
            df[col] = df[col].apply(converter_mes_envio)

    return df


def pegar_coluna_existente(df, opcoes, nome_logico):
    for col in opcoes:
        if col in df.columns:
            return col

    raise KeyError(
        f"Nenhuma coluna encontrada para {nome_logico}. "
        f"Procurei por: {opcoes}. "
        f"Colunas disponíveis: {list(df.columns)}"
    )


def aplicar_formatacao_datas_excel(caminho_excel):
    wb = load_workbook(caminho_excel)
    ws = wb.active

    headers = {
        cell.value: cell.column
        for cell in ws[1]
        if cell.value is not None
    }

    formatos = {
        "DATA_DE_ENVIO": "DD/MM/YYYY",
        "MES_DE_ENVIO": "MM/YYYY",
        "Data de Envio": "DD/MM/YYYY",
        "Mês de Envio": "MM/YYYY",
    }

    for nome_coluna, formato in formatos.items():
        if nome_coluna in headers:
            col_idx = headers[nome_coluna]
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col_idx).number_format = formato

    wb.save(caminho_excel)


def get_conf():
    path = os.path.join(PROJECT_ROOT, "conf", "conf_cusd_invoices.yml")
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def get_query(query_name, model_parameters):
    path = os.path.join(PROJECT_ROOT, "sql", model_parameters["queries"][query_name])
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


# ============================================================
# DEFINIÇÃO DA PASTA DA DISTRIBUIDORA ESCOLHIDA
# ============================================================

if distribuidora_escolhida not in LEITORES:
    raise ValueError(
        f"Distribuidora sem leitor configurado: {distribuidora_escolhida}. "
        f"Disponíveis: {list(LEITORES)}"
    )

nome_pasta = pastas_especiais.get(distribuidora_escolhida, distribuidora_escolhida)
folder_rateio = os.path.join(base_rateio, nome_pasta)
folder_rateio_mes = os.path.join(folder_rateio, mes_pasta_rateio)

if not Path(folder_rateio_mes).exists():
    raise FileNotFoundError(
        f"Pasta do mês não encontrada: {folder_rateio_mes}. "
        f"Verifique o mes_pasta_rateio."
    )

print("Hist existe:", Path(folder_hist).exists())
print("Pasta rateio existe:", Path(folder_rateio_mes).exists())
print("Distribuidora escolhida:", distribuidora_escolhida)
print("Pasta usada:", folder_rateio_mes)


print(f"{Fore.MAGENTA}\n=== ETAPA 2: Conexão e leitura das bases ==={Style.RESET_ALL}")

conf = get_conf()
model_parameters = conf["model_parameters"]
access, conn = initialize_connections()

df_usinas = pd.read_sql(get_query("hist_usina_livia", model_parameters), conn)
df_crm = pd.read_sql(get_query("crm", model_parameters), conn)


print(f"{Fore.MAGENTA}\n=== ETAPA 3: Preparação das bases auxiliares ==={Style.RESET_ALL}")

df_usinas_sel = df_usinas[[
    "DISTRIBUIDORA",
    "NUMERO_DA_INSTALACAO",
    "PROJETOS",
    "ACRONIMO"
]].copy()

df_crm_sel = df_crm[[
    "DISTRIBUIDORA",
    "UNIDADE_CONSUMIDORA",
    "NUMERO_CONTA_CONTRATO"
]].copy()


map_distribuidoras = {
    "CELGD": "EQUATORIAL GO",
    "EQUATORIAL GO": "EQUATORIAL GO",
    "COELBA": "COELBA",
    "CEMIGD": "CEMIG",
    "CEMIG": "CEMIG",
    "CEPISA": "EQUATORIAL PI",
    "EQUATORIAL PI": "EQUATORIAL PI",
    "CEMAR": "EQUATORIAL MA",
    "EQUATORIAL MA": "EQUATORIAL MA",
    "CELPA": "EQUATORIAL PA",
    "EQUATORIAL PA": "EQUATORIAL PA",
    "CPFLPAULISTA": "CPFL PAULISTA",
    "CPFL PAULISTA": "CPFL PAULISTA",
    "ENELRJ": "ENEL RJ",
    "ENEL RJ": "ENEL RJ",
    "ENELCE": "ENEL CE",
    "ENEL CE": "ENEL CE",
    "EMT": "EMT",
    "ENERGISA MT": "EMT",
    "CELPE": "CELPE"
}

for df_aux in [df_usinas_sel, df_crm_sel]:
    df_aux["DISTRIBUIDORA"] = (
        df_aux["DISTRIBUIDORA"]
        .astype(str)
        .str.strip()
        .str.upper()
        .map(map_distribuidoras)
        .fillna(df_aux["DISTRIBUIDORA"].astype(str).str.strip().str.upper())
    )


df_usinas_sel = df_usinas_sel.rename(columns={
    "DISTRIBUIDORA": "Distribuidora",
    "NUMERO_DA_INSTALACAO": "UG",
    "PROJETOS": "Usina",
    "ACRONIMO": "Acrônimo",
})

df_crm_sel = df_crm_sel.rename(columns={
    "DISTRIBUIDORA": "Distribuidora",
    "UNIDADE_CONSUMIDORA": "UC",
    "NUMERO_CONTA_CONTRATO": "CC",
})


print(f"{Fore.MAGENTA}\n=== ETAPA 4: Leitura do novo rateio {distribuidora_escolhida} ==={Style.RESET_ALL}")

# Leitura + trava de qualidade (o rateio precisa fechar ~100% por UG).
# Erros aqui são acionáveis pelo usuário — arquivo ainda não baixado do
# OneDrive, rateio que não fecha 100%, nome de arquivo fora do padrão, etc. —
# então mostramos a mensagem limpa e encerramos, sem despejar traceback.
try:
    df_novo = LEITORES[distribuidora_escolhida](folder_rateio_mes, data_hoje, mes_envio)
    validar_soma_rateio(df_novo, tolerancia=0.01, bloquear=True)
except (FileNotFoundError, ValueError) as e:
    print(f"{Fore.RED}\n⛔ Processo interrompido:{Style.RESET_ALL}\n{e}")
    sys.exit(1)

df_novo = padronizar_datas_rateio(df_novo)


print(f"{Fore.MAGENTA}\n=== ETAPA 5: Padronização das chaves ==={Style.RESET_ALL}")

df_novo = padronizar_chaves_para_merge(df_novo)
df_usinas_sel = padronizar_chaves_para_merge(df_usinas_sel)
df_crm_sel = padronizar_chaves_para_merge(df_crm_sel)


print(f"{Fore.YELLOW}\n=== DEBUG: comparação de UGs ==={Style.RESET_ALL}")

ugs_df_novo = sorted(df_novo["UG"].dropna().unique())
print(f"\nUGs em df_novo ({len(ugs_df_novo)}):")
print(ugs_df_novo)

ugs_usinas_base = sorted(
    df_usinas_sel.loc[
        df_usinas_sel["Distribuidora"] == distribuidora_escolhida,
        "UG"
    ].dropna().unique()
)

print(f"\nUGs em df_usinas_sel para {distribuidora_escolhida} ({len(ugs_usinas_base)}):")
print(ugs_usinas_base)

set_novo = set(ugs_df_novo)
set_usinas = set(ugs_usinas_base)

print(f"\nUGs no df_novo que NÃO existem em df_usinas_sel:")
print(sorted(set_novo - set_usinas))

print(f"\nUGs em df_usinas_sel que NÃO aparecem no df_novo:")
print(sorted(set_usinas - set_novo))

print(f"\nUGs em comum: {len(set_novo & set_usinas)}")


print("\n=== DIAGNÓSTICO DE DUPLICADOS ANTES DOS MERGES ===")

dup_usinas = (
    df_usinas_sel
    .groupby(["Distribuidora", "UG"], dropna=False)
    .size()
    .reset_index(name="qtd")
    .query("qtd > 1")
)

print("\nDuplicados em df_usinas_sel por Distribuidora + UG:")
print(dup_usinas)

dup_crm = (
    df_crm_sel
    .groupby(["Distribuidora", "UC"], dropna=False)
    .size()
    .reset_index(name="qtd")
    .query("qtd > 1")
)

print("\nDuplicados em df_crm_sel por Distribuidora + UC:")
print(dup_crm)


df_usinas_sel = (
    df_usinas_sel
    .sort_values(["Distribuidora", "UG", "Usina", "Acrônimo"])
    .drop_duplicates(subset=["Distribuidora", "UG"], keep="first")
)

df_crm_sel = (
    df_crm_sel
    .sort_values(["Distribuidora", "UC"])
    .drop_duplicates(subset=["Distribuidora", "UC"], keep="first")
)

print("\nApós remover duplicados:")
print("Linhas df_usinas_sel:", len(df_usinas_sel))
print("Linhas df_crm_sel:", len(df_crm_sel))


print(f"{Fore.MAGENTA}\n=== ETAPA 6: Enriquecendo df_novo com dados da usina ==={Style.RESET_ALL}")

linhas_antes_usina = len(df_novo)

df_novo = df_novo.merge(
    df_usinas_sel[[
        "Distribuidora",
        "UG",
        "Usina",
        "Acrônimo"
    ]],
    on=["Distribuidora", "UG"],
    how="left",
    validate="many_to_one"
)

df_novo["CC/UG"] = pd.NA

print("Linhas antes do merge usina:", linhas_antes_usina)
print("Linhas depois do merge usina:", len(df_novo))

print(f"\n{Fore.YELLOW}DEBUG pós-merge usina:{Style.RESET_ALL}")
print("Linhas com Usina vazia:", df_novo["Usina"].isna().sum())

chaves_sem_match = (
    df_novo[df_novo["Usina"].isna()][["Distribuidora", "UG"]]
    .drop_duplicates()
)

print("Chaves Distribuidora + UG sem match:")
print(chaves_sem_match)


print(f"{Fore.MAGENTA}\n=== ETAPA 7: Enriquecendo df_novo com CC do CRM ==={Style.RESET_ALL}")

linhas_antes_crm = len(df_novo)

df_novo = df_novo.merge(
    df_crm_sel[["Distribuidora", "UC", "CC"]],
    on=["Distribuidora", "UC"],
    how="left",
    validate="many_to_one"
)

print("Linhas antes do merge CRM:", linhas_antes_crm)
print("Linhas depois do merge CRM:", len(df_novo))


print(f"\n{Fore.YELLOW}DEBUG pós-merge CRM:{Style.RESET_ALL}")

ucs_sem_cc = df_novo[df_novo["CC"].isna()][["Distribuidora", "UG", "UC", "%"]].copy()

print(f"Linhas sem CC: {len(ucs_sem_cc)}")
print(f"UCs distintas sem match no CRM: {ucs_sem_cc['UC'].nunique()}")
print("\nDetalhe:")
print(ucs_sem_cc.to_string(index=False))

saida_ucs = os.path.join(
    folder_hist,
    f"ucs_sem_cc_{distribuidora_escolhida.replace(' ', '_')}_{datetime.today().strftime('%Y%m%d_%H%M%S')}.csv"
)

ucs_sem_cc.to_csv(saida_ucs, index=False, encoding="utf-8-sig")
print(f"\nArquivo de UCs sem CC salvo: {saida_ucs}")


print(f"{Fore.MAGENTA}\n=== ETAPA 8: Validação da df_novo antes do histórico ==={Style.RESET_ALL}")

print("\nCampos principais antes do histórico:")
print(df_novo[[
    "Distribuidora",
    "UG",
    "Usina",
    "CC/UG",
    "Acrônimo",
    "UC",
    "CC",
    "%",
    "Data de Envio",
    "Mês de Envio"
]].head(20))

print("\nValores não preenchidos antes do histórico:")
print("Usina vazia:", df_novo["Usina"].isna().sum())
print("CC/UG vazio:", df_novo["CC/UG"].isna().sum())
print("Acrônimo vazio:", df_novo["Acrônimo"].isna().sum())
print("CC vazio:", df_novo["CC"].isna().sum())

print("\nSoma de % por UG na df_novo:")
print(df_novo.groupby("UG")["%"].sum())


print(f"{Fore.MAGENTA}\n=== ETAPA 9: Leitura da HistRateio pronta ==={Style.RESET_ALL}")

caminho_hist = get_latest_excel_file(folder_hist, allocation_month)

df_histraterio = pd.read_excel(caminho_hist)
df_histraterio = padronizar_chaves_para_merge(df_histraterio)
df_histraterio = padronizar_datas_rateio(df_histraterio)

print("Linhas df_histraterio:", len(df_histraterio))
print("Linhas df_novo:", len(df_novo))


print(f"{Fore.MAGENTA}\n=== ETAPA 9.1: Ajuste de nomes do df_novo para padrão do histórico ==={Style.RESET_ALL}")

map_colunas_df_novo = {
    "Distribuidora": "DISTRIBUIDORA",
    "Usina": "USINA",
    "Acrônimo": "ACRONIMO",
    "UG": "UG",
    "CC/UG": "CC/UG",
    "UC": "UC",
    "CC": "CC",
    "%": "%",
    "Data de Envio": "DATA_DE_ENVIO",
    "Mês de Envio": "MES_DE_ENVIO"
}

rename_final = {
    col_origem: col_destino
    for col_origem, col_destino in map_colunas_df_novo.items()
    if col_origem in df_novo.columns and col_destino in df_histraterio.columns
}

df_novo = df_novo.rename(columns=rename_final)
df_novo = padronizar_datas_rateio(df_novo)

col_mes_envio = pegar_coluna_existente(
    df_novo,
    ["MES_DE_ENVIO", "Mês de Envio"],
    "mês de envio"
)

col_acronimo = pegar_coluna_existente(
    df_novo,
    ["ACRONIMO", "Acrônimo"],
    "acrônimo"
)

print("Coluna usada para mês de envio:", col_mes_envio)
print("Coluna usada para acrônimo:", col_acronimo)


# ============================================================
# CAMPOS PADRÃO DO HISTÓRICO
# ============================================================

if "UC/CC" in df_histraterio.columns:
    df_novo["UC/CC"] = df_novo["UC"]

if "Id" in df_histraterio.columns:
    df_novo["Id"] = (
        df_novo["UG"].astype(str)
        + " - "
        + df_novo["UC"].astype(str)
        + " - "
        + df_novo[col_mes_envio]
            .dt.strftime("%m/%Y")
            .str.replace(r"^0", "", regex=True)
    )

if "Chave" in df_histraterio.columns:
    df_novo["Chave"] = (
        df_novo[col_acronimo].astype(str)
        + "-"
        + df_novo[col_mes_envio]
            .dt.strftime("%m/%Y")
            .str.replace(r"^0", "", regex=True)
    )

if "Aceitou Rateio?" in df_histraterio.columns:
    df_novo["Aceitou Rateio?"] = "NÃO"


print("\nColunas renomeadas no df_novo:")
print(rename_final)

print("\nValidação após renomear:")
for col in ["DISTRIBUIDORA", "Distribuidora", "USINA", "Usina", "ACRONIMO", "Acrônimo", "UG", "CC/UG", "UC", "CC"]:
    if col in df_novo.columns:
        print(f"{col} vazio:", df_novo[col].isna().sum())


print(f"{Fore.MAGENTA}\n=== ETAPA 10: Ajuste de colunas antes do concat ==={Style.RESET_ALL}")

for col in df_histraterio.columns:
    if col not in df_novo.columns:
        df_novo[col] = pd.NA

df_novo = df_novo[df_histraterio.columns]
df_novo = padronizar_datas_rateio(df_novo)


print(f"{Fore.MAGENTA}\n=== ETAPA 10.1: Validação final do df_novo já no layout do histórico ==={Style.RESET_ALL}")

colunas_validacao = [
    col for col in [
        "DISTRIBUIDORA",
        "Distribuidora",
        "USINA",
        "Usina",
        "ACRONIMO",
        "Acrônimo",
        "UG",
        "CC/UG",
        "UC/CC",
        "UC",
        "CC",
        "%",
        "DATA_DE_ENVIO",
        "Data de Envio",
        "MES_DE_ENVIO",
        "Mês de Envio",
        "Aceitou Rateio?"
    ]
    if col in df_novo.columns
]

print(df_novo[colunas_validacao].head(30))

for col in ["USINA", "Usina", "ACRONIMO", "Acrônimo", "CC/UG", "CC"]:
    if col in df_novo.columns:
        print(f"{col} vazio após ajuste final:", df_novo[col].isna().sum())

if "%" in df_novo.columns:
    print("\nSoma de % por UG no layout final:")
    print(df_novo.groupby("UG")["%"].sum())


print(f"{Fore.MAGENTA}\n=== ETAPA 11: Concat final ==={Style.RESET_ALL}")

df_final = pd.concat(
    [df_histraterio, df_novo],
    ignore_index=True
)

df_final = padronizar_datas_rateio(df_final)

print("Linhas histórico:", len(df_histraterio))
print("Linhas novo:", len(df_novo))
print("Linhas final:", len(df_final))

if len(df_final) == len(df_histraterio) + len(df_novo):
    print("OK: concat realizado corretamente.")
else:
    print("ATENÇÃO: quantidade de linhas inesperada no concat.")


print(f"{Fore.MAGENTA}\n=== ETAPA 12: Exportação ==={Style.RESET_ALL}")

saida = os.path.join(
    folder_hist,
    f"HistRateio_merge_{distribuidora_escolhida.replace(' ', '_')}_{datetime.today().strftime('%Y%m%d_%H%M%S')}.xlsx"
)

df_final.to_excel(saida, index=False)
aplicar_formatacao_datas_excel(saida)

print(f"\nArquivo salvo: {saida}")